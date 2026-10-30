# -*- coding: utf-8 -*-
"""生成最终报告：第一至七章 + 附录（模板底稿）"""
from docx import Document

from make_report_appendix import add_appendix
from make_report_ch5 import ch5_cang, ch5_ren
from make_report_ch67 import ch67_cang, ch67_ren
from make_report_full import OUT, PEOPLE, build
from report_extras import (add_page_numbers, add_toc, fix_cover, fix_toc_styles,
                            normalize_cover, disable_auto_update)
from report_shots import add_screenshots
from report_base import base_doc


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for name, sid in PEOPLE.items():
        ch5 = ch5_ren if name == "任绪锋" else ch5_cang
        ch67 = ch67_ren if name == "任绪锋" else ch67_cang
        doc = base_doc(name, sid)
        build(name, sid, ch5_fn=ch5, doc=doc, ch67_fn=ch67, appx_fn=add_appendix)
        add_screenshots(doc, name)
        add_toc(doc)
        fix_toc_styles(doc)
        add_page_numbers(doc)
        fix_cover(doc)
        normalize_cover(doc)
        disable_auto_update(doc)
        path = OUT / f"报告-最终-{name}-{sid}.docx"
        try:
            doc.save(str(path))
        except PermissionError:
            print(f"[跳过] {path.name} 正被 Word 占用")
            continue
        d = Document(str(path))
        chars = sum(len(x.text) for x in d.paragraphs)
        print(f"[OK] {path.name}：{len(d.paragraphs)} 段｜正文 {chars} 字｜表 {len(d.tables)}｜图 {len(d.inline_shapes)}｜{path.stat().st_size} bytes")


if __name__ == "__main__":
    main()
