# -*- coding: utf-8 -*-
"""共用：以课程模板为底稿

作用：
- 打开模板 docx
- 删掉模板里的示例章节与提示段
- 填封面表格（题目/姓名/学号）与「项目组成员及工作分配」表
- 返回一个可直接续写正文的 Document

三种报告（骨架 / 前四章 / 前五章）都从这里取底稿，保证封面与前置部分完全一致。
"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor

TPL = Path(r"D:\vrptw\docs\课程模板\报告撰写模板-2026修订后.docx")
TOPIC = "基于改进蚁群算法的带时间窗配送路径优化系统设计与实现"
INFO_FILE = Path("D:/vrptw/docs/封面信息.json")


def load_cover_info() -> dict:
    """封面信息（班级/指导教师/职称/成员）从 JSON 读，改一处全文档生效"""
    if INFO_FILE.exists():
        return json.loads(INFO_FILE.read_text(encoding="utf-8"))
    return {}

ROLE_REN = "负责算例数据与算法引擎（M1 算例与数据管理、M2 算法引擎、M5 约束与场景配置）"
ROLE_CANG = "负责可视化界面与实验统计（M3 路线可视化、M4 实验管理与统计、M6 方案对比与报表）"


def _text(el) -> str:
    return "".join(n.text or "" for n in el.iter(qn("w:t"))).strip()


def set_cell(cell, text: str) -> None:
    """只改第一个 run（保住字体字号），并把颜色重置为黑色——模板示例文字是红色。"""
    par = cell.paragraphs[0]
    if not par.runs:
        r0 = par.add_run(text)
        r0.font.color.rgb = RGBColor(0, 0, 0)
        return
    par.runs[0].text = text
    par.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    for r in par.runs[1:]:
        r._element.getparent().remove(r._element)


def prepare(doc, name: str, sid: str):
    body = doc.element.body

    # 1) 删「注：请根据自己组…」提示段
    for child in list(body.iterchildren()):
        if child.tag == qn("w:p") and _text(child).startswith("注：请根据自己组"):
            body.remove(child)

    # 2) 删示例章节：从「第一章」到 sectPr 之前
    start = None
    for child in body.iterchildren():
        if child.tag == qn("w:p") and _text(child).startswith("第一章"):
            start = child
            break
    if start is None:
        raise SystemExit("模板结构变了：没找到「第一章」")
    el = start
    while el is not None and el.tag != qn("w:sectPr"):
        nxt = el.getnext()
        body.remove(el)
        el = nxt

    # 3) 封面表格
    t = doc.tables[0]
    set_cell(t.rows[0].cells[1], TOPIC)
    set_cell(t.rows[2].cells[1], name)
    set_cell(t.rows[2].cells[3], sid)
    info = load_cover_info()
    set_cell(t.rows[3].cells[3], info.get("banji") or "【待填】")
    set_cell(t.rows[4].cells[1], info.get("advisor") or "【待填】")
    set_cell(t.rows[4].cells[3], info.get("title") or "【待填】")

    # 4) 项目组成员及工作分配
    m = doc.tables[1]
    rows = info.get("members") or [
        ["组长", "任绪锋", "2318140115", ROLE_REN],
        ["组员", "苍韵同", "2318140111", ROLE_CANG],
    ]
    for i, (role, nm, sid_, desc) in enumerate(rows[:2]):
        set_cell(m.rows[i].cells[0], role)
        set_cell(m.rows[i].cells[1], nm)
        set_cell(m.rows[i].cells[2], sid_)
        set_cell(m.rows[i].cells[3], desc)
    return doc


def base_doc(name: str, sid: str) -> Document:
    """模板底稿：封面/课程任务要求/成员表就绪，示例章节已清空"""
    return prepare(Document(str(TPL)), name, sid)
