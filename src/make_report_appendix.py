# -*- coding: utf-8 -*-
"""附录：A 核心代码清单 / B 运行效果图（补页数，也为答辩留证据）"""
from pathlib import Path

from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from make_report_ch5 import code
from make_report_full import h, p

VIZ = Path("D:/vrptw/results/viz")


def fig(doc, path: Path, caption: str, width_cm: float = 15.0) -> None:
    if not path.exists():
        p(doc, f"（缺图：{path.name}）", size=10)
        return
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.add_run().add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.font.size = __import__("docx").shared.Pt(10.5)


def add_appendix(doc) -> None:
    h(doc, "附录 A  核心代码清单", 1)
    p(doc, "本附录摘录各模块的关键实现，便于说明设计思路与复现实验。完整代码见项目目录 "
           "D:/vrptw（18 个 Python 文件，约 3300 行）。")

    h(doc, "A.1  算例解析（instance.py）", 2)
    p(doc, "用“整行是否全为数字”区分表头与数据行，按 VEHICLE / CUSTOMER 两段做状态机解析。")
    code(doc, '''
section = None
for raw in text.splitlines():
    line = raw.strip()
    if not line:
        continue
    key = line.upper()
    if key.startswith("VEHICLE"):
        section = "vehicle";  continue
    if key.startswith("CUSTOMER"):
        section = "customer"; continue
    nums = _numeric_tokens(line)          # 整行都是数字才返回
    if not nums:
        continue
    if section == "vehicle" and vehicle_count is None:
        vehicle_count, capacity = int(nums[0]), float(nums[1])
    elif section == "customer" and len(nums) >= 7:
        customers.append(Customer(int(nums[0]), nums[1], nums[2],
                                  nums[3], nums[4], nums[5], nums[6]))
''')

    h(doc, "A.2  可行解构造（aco.py）", 2)
    p(doc, "候选客户先在构造阶段过滤掉“装不下 / 赶不上时间窗”的，再做轮盘赌选择；"
           "这样不会产生大量不可行解，也不需要事后修复。")
    code(doc, '''
cands = []
for cid in unvisited:
    c = inst.customers[cid]
    if load + c.demand > inst.capacity:
        continue
    trav = inst.time(prev, cid)
    start = max(t + trav, c.ready)
    if start > c.due:
        continue
    cands.append((cid, inst.dist(prev, cid), trav,
                  start - (t + trav), c.due - start))
if not cands:
    break
weights = [(tau[prev][cid] ** alpha) * (self._eta(d, wait, slack) ** beta)
           for (cid, d, trav, wait, slack) in cands]
cid, d, trav, wait, slack = self._roulette(cands, weights)
''')

    h(doc, "A.3  信息素上下界（aco.py，MMAS）", 2)
    p(doc, "每次更新后把信息素夹在 [tau_min, tau_max] 内，抑制某条路径过早垄断，缓解早熟。")
    code(doc, '''
tau_max = 1.0 / (rho * best_cost)
tau_min = tau_max / (2.0 * self.n)
for i in range(self.n):
    for j in range(self.n):
        if i != j:
            tau[i][j] = min(tau_max, max(tau_min, tau[i][j]))
''')

    h(doc, "A.4  并行实验调度（run_experiments.py）", 2)
    p(doc, "把实验拆成最小任务交给进程池；主进程边取结果边写 CSV 并立即刷盘，中断不丢数据。")
    code(doc, '''
tasks = [(str(f), m, r) for f in files for m in methods for r in range(runs)]
with mp.Pool(jobs) as pool:
    for done, row in enumerate(pool.imap_unordered(_work, tasks), 1):
        rows.append(row)
        writer.writerow({k: row.get(k) for k in fieldnames})
        fh.flush()
        if done % 5 == 0:
            print(f"进度 {done}/{len(tasks)}")
''')

    h(doc, "A.5  实验记录落库（store.py）", 2)
    p(doc, "同一套 SQL 同时兼容 SQLite 与 MySQL：占位符统一写问号，由 db.py 在 MySQL 后端转成百分号形式。")
    code(doc, '''
def insert_experiment(db, instance_id, method, run_index, metrics,
                      seed=None, iterations=None, seconds=None):
    cur = db.execute(
        "INSERT INTO experiment (instance_id, method, run_index, seed,"
        " iterations, distance, vehicles, feasible, on_time_rate, seconds)"
        " VALUES (?,?,?,?,?,?,?,?,?,?)",
        (instance_id, method, run_index, seed, iterations,
         metrics.get("total_distance"), metrics.get("vehicles"),
         1 if metrics.get("feasible") else 0,
         metrics.get("on_time_rate"), seconds))
    db.commit()
    return cur.lastrowid
''')

    h(doc, "附录 B  运行效果图", 1)
    p(doc, "以下图由项目脚本直接生成（python viz.py data/solomon/100/c101.txt --iters=30）。")
    fig(doc, VIZ / "C101_before_after.png", "图 B-1  C101 优化前后路线对比（左：最近邻 1870.73/21 车；右：改进版 829.01/10 车）", 15.5)
    fig(doc, VIZ / "C101_convergence.png", "图 B-2  收敛曲线（每代最优里程）", 14.0)
    fig(doc, VIZ / "C101_ablation.png", "图 B-3  消融对比（各方法总里程，柱上标注与最优解的差距）", 14.0)
