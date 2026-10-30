# -*- coding: utf-8 -*-
"""以课程模板为底稿生成「完整版」报告（第一至五章）

产出：D:\\vrptw\\docs\\报告-完整-<姓名>-<学号>.docx
"""
from __future__ import annotations

from docx import Document

from make_report_ch5 import ch5_cang, ch5_ren
from make_report_full import OUT, PEOPLE, build
from report_base import base_doc


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for name, sid in PEOPLE.items():
        fn = ch5_ren if name == "任绪锋" else ch5_cang
        doc = base_doc(name, sid)
        build(name, sid, ch5_fn=fn, doc=doc)

        path = OUT / f"报告-完整-{name}-{sid}.docx"
        try:
            doc.save(str(path))
        except PermissionError:
            print(f"[跳过] {path.name} 正被 Word 占用")
            continue

        d = Document(str(path))
        chars = sum(len(x.text) for x in d.paragraphs)
        t = d.tables[0]
        print(f"[OK] {path.name}：{len(d.paragraphs)} 段｜{chars} 字｜"
              f"图 {len(d.inline_shapes)}｜表 {len(d.tables)}｜{path.stat().st_size} bytes")
        print(f"      封面：{t.rows[2].cells[1].text.strip()} / {t.rows[2].cells[3].text.strip()}"
              f"｜成员：{d.tables[1].rows[0].cells[1].text.strip()}、"
              f"{d.tables[1].rows[1].cells[1].text.strip()}")


if __name__ == "__main__":
    main()
