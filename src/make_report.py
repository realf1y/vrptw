# -*- coding: utf-8 -*-
"""生成报告骨架（每人一份 docx），按课程模板的七章结构

产出：D:\\vrptw\\docs\\报告骨架-任绪锋-2318140115.docx
      D:\\vrptw\\docs\\报告骨架-苍韵同-2318140111.docx

说明：第一至四章全组共用；第五章、第六章按各自负责模块不同；第七章附 12 周周记。
      正文里的「写作要点」是给你们的提示，定稿前要删掉替换成真实内容。
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
OUT = Path("D:/vrptw/docs")

TOPIC = "基于改进蚁群算法的带时间窗配送路径优化系统设计与实现"

PEOPLE = {
    "任绪锋": {"sid": "2318140115",
               "modules": ["算例与数据管理模块", "蚁群算法引擎模块", "约束与场景配置模块"]},
    "苍韵同": {"sid": "2318140111",
               "modules": ["路线可视化与交互模块", "实验管理与统计模块", "方案对比与报表模块"]},
}


def style_run(run, size=12, bold=False, cn="宋体"):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn)


def para(doc, text, size=12, bold=False, indent=True, align=None, cn="宋体",
         space_after=6, color=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(22)
    pf.space_after = Pt(space_after)
    if indent:
        pf.first_line_indent = Pt(24)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    style_run(r, size=size, bold=bold, cn=cn)
    if color is not None:
        r.font.color.rgb = color
    return p


def heading(doc, text, level=1):
    h = doc.add_heading(level=level)
    r = h.add_run(text)
    style_run(r, size=16 if level == 1 else 14, bold=True, cn="黑体")
    return h


def tip(doc, text):
    """写作提示：灰色小字，定稿前删掉"""
    from docx.shared import RGBColor
    return para(doc, "写作要点：" + text, size=10, indent=True,
                color=RGBColor(0x99, 0x99, 0x99))


def build(name: str, doc=None) -> Document:
    info = PEOPLE[name]
    fresh = doc is None
    if fresh:
        doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    if fresh:
        # ---------------- 封面 ----------------
        para(doc, "《软件工程系统实训》综合报告", size=22, bold=True, indent=False,
             align=WD_ALIGN_PARAGRAPH.CENTER, cn="黑体", space_after=18)
        cover = [
            ("题　　目：", TOPIC),
            ("起讫日期：", "2026 年 9 月 21 日 至 2026 年 12 月 20 日"),
            ("学生姓名：", name),
            ("学　　号：", info["sid"]),
            ("所在院系：", "轨道智能工程学院"),
            ("专业班级：", "【待填】"),
            ("指导教师：", "【待填】"),
        ]
        for k, v in cover:
            para(doc, f"{k}{v}", size=14, indent=False, space_after=8)
        doc.add_page_break()

        # ---------------- 课程任务及要求 ----------------
        heading(doc, "课程任务及要求", 1)
        tip(doc, "本部分按课程发放的模板原文保留（实践目的与要求、内容及步骤、进度安排、答辩与成绩评定）。")

    # ---------------- 第一章 ----------------
    heading(doc, "第一章  系统概述", 1)
    heading(doc, "1.1  应用领域调研", 2)
    para(doc, "配送路径优化属于物流与供应链领域的核心调度问题。本节选取三类同类系统做对比："
              "① 国外商用路线规划平台（如 Route4Me）；② 开源求解器与算法库（如 Google OR-Tools）；"
              "③ 国内即时配送平台的调度系统。")
    tip(doc, "补 3 个真实产品的名称、官网与功能截图；用一两句话说明本系统与它们的相似处（都做路线优化）"
             "和不同点（本系统面向教学与实验，算法可替换、指标可复现、数据可导出）。参考文献处标 [1]。")
    heading(doc, "1.2  选题合理性", 2)
    para(doc, "带时间窗的车辆路径问题（VRPTW）是 NP-hard 组合优化问题：客户数量增加时可行解数量呈指数增长，"
              "人工排线既慢又难保证质量。同时在真实配送中，客户对送达时间有明确要求，"
              "车辆有载重上限，二者共同构成强约束。因此研究并实现一个可配置、可复现的求解系统具有实际意义。")
    heading(doc, "1.3  系统概述", 2)
    para(doc, "本系统以改进蚁群算法为核心，提供算例管理、场景配置、算法求解、路线可视化、"
              "实验统计与报表导出六项能力，支持在标准算例上批量实验并对比不同算法配置的效果。")

    # ---------------- 第二章 ----------------
    heading(doc, "第二章  可行性研究及需求分析", 1)
    heading(doc, "2.1  可行性研究", 2)
    para(doc, "技术可行性：算法与系统均基于 Python 标准库实现，仅数据库驱动使用第三方包，普通笔记本即可运行；"
              "经济可行性：无需购买商用求解器许可；操作可行性：命令行与网页两种入口，学习成本低；"
              "法律可行性：使用公开标准算例，不涉及真实客户隐私数据。")
    heading(doc, "2.2  系统需求概述", 2)
    para(doc, "系统参与者为实验人员，主要交互包括：选择算例、配置业务场景与算法参数、启动求解、"
              "查看路线与指标、对比多方案、导出报表。系统总用例图见图 2-1。")
    para(doc, "图 2-1  系统用例图（插入 results/diagrams/usecase.png）", indent=False,
         align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
    heading(doc, "2.3  运行优化算法用例分析", 2)
    heading(doc, "2.3.1  用例描述", 3)
    para(doc, "参与者选择算例与场景后启动求解，系统校验载重与时间窗约束，按设定代数迭代，"
              "输出最优路线与里程、车辆数、准时率等指标。前置条件：算例文件存在且格式合法。"
              "后置条件：生成一条实验记录并落库。")
    heading(doc, "2.3.2  用例对应的活动图", 3)
    para(doc, "图 2-2  运行优化算法活动图（待绘制：选择算例 → 校验数据 → 构造初始解 → 迭代改进 → 输出结果）",
         indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
    heading(doc, "2.4  配置业务场景用例分析", 2)
    heading(doc, "2.4.1  用例描述", 3)
    para(doc, "参与者设置载重上限、车辆数上限、车速、时间窗模式（硬/软）与惩罚系数、客户优先级，"
              "系统将其保存为场景配置，供后续实验复用。")
    heading(doc, "2.4.2  用例对应的活动图", 3)
    para(doc, "图 2-3  配置业务场景活动图（待绘制：新建场景 → 填写参数 → 校验取值范围 → 保存 JSON）",
         indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=10)

    # ---------------- 第三章 ----------------
    heading(doc, "第三章  项目系统的数据库设计", 1)
    heading(doc, "3.1  概念设计", 2)
    para(doc, "系统包含算例、场景、实验记录、方案路线四类核心实体：一个算例可对应多次实验，"
              "一个场景也可对应多次实验，一次实验产出多条路线。E-R 图见图 3-1。")
    para(doc, "图 3-1  系统数据库 E-R 图（插入 results/diagrams/er.png）", indent=False,
         align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
    heading(doc, "3.2  逻辑设计", 2)
    para(doc, "将 E-R 图转换为关系模式（下划线为主键）：")
    for line in [
        "算例 instance（id，name，size，n_customers，capacity，vehicle_count，speed，total_demand，source_path，created_at）",
        "场景 scene（id，name，config_json，note，created_at）",
        "实验记录 experiment（id，instance_id，scene_id，method，run_index，seed，iterations，distance，vehicles，feasible，tw_violation，on_time_rate，load_rate，seconds，created_at）",
        "方案路线 solution（id，experiment_id，route_index，customers，stop_count，route_length，route_load）",
    ]:
        para(doc, line, size=11, indent=True, space_after=2)
    heading(doc, "3.3  物理设计", 2)
    para(doc, "采用 MySQL 8/9（InnoDB，utf8mb4）。主要设计要点：")
    for line in [
        "① 主键统一为 INT AUTO_INCREMENT；",
        "② instance.name、scene.name 建唯一索引，避免重复导入；",
        "③ experiment 对 instance_id、method 建普通索引，支撑按算例/方法的聚合查询；",
        "④ 外键 fk_exp_instance、fk_exp_scene、fk_sol_exp 均设 ON DELETE CASCADE（solution 随实验级联删除）；",
        "⑤ 建立视图 kpi_view 预聚合「算例 × 方法」的平均里程、最优里程、车辆数、准时率、装载率与耗时。",
    ]:
        para(doc, line, size=11, space_after=2)
    para(doc, "完整建表语句由 python store.py schema --mysql 导出，可直接执行。", size=11)
    heading(doc, "3.4  数据库连接", 2)
    para(doc, "系统通过 pymysql 连接 MySQL，连接参数集中在 mysql_config.json（host、port、user、password、database、charset），"
              "代码侧由 db.py 统一封装，同时兼容 SQLite（本地零配置开发）。"
              "生产环境使用独立业务账号 vrptw（仅授予 vrptw 库权限），不直接使用 root。")

    # ---------------- 第四章 ----------------
    heading(doc, "第四章  项目系统的总体设计", 1)
    heading(doc, "4.1  系统总体架构设计", 2)
    para(doc, "系统分为四层：数据层（算例解析、距离矩阵、SQLite/MySQL 持久化）、"
              "算法层（蚁群构造解、局部搜索、信息素更新）、服务层（实验调度、指标计算、报表导出）、"
              "展示层（路线图、收敛曲线、消融对比与 KPI 报表）。层间通过函数接口与数据表解耦。")
    heading(doc, "4.2  系统界面设计", 2)
    para(doc, "界面分为三屏：实验配置、路线对比、实验报表，原型见图 4-1。", indent=True)
    para(doc, "图 4-1  系统界面原型（插入 results/diagrams/ui.png）", indent=False,
         align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
    heading(doc, "4.3  系统功能模块设计", 2)
    para(doc, "系统共 6 个功能模块：M1 算例与数据管理、M2 算法引擎、M3 路线可视化与交互、"
              "M4 实验管理与统计、M5 约束与场景配置、M6 方案对比与报表。各模块职责与接口见下节。")

    # ---------------- 第五章 ----------------
    heading(doc, "第五章  项目系统的实现", 1)
    for i, mod in enumerate(info["modules"], start=1):
        heading(doc, f"5.{i}  {mod}的设计与实现", 2)
        para(doc, f"（本模块由 {name} 负责，需写：功能需求、类/函数设计、关键代码说明、运行截图。）")
        tip(doc, "建议按「输入—处理—输出」三段写；附 1 张流程图或类图 + 1 段核心代码 + 1 张运行效果截图。")

    # ---------------- 第六章 ----------------
    heading(doc, "第六章  项目系统的测试", 1)
    heading(doc, "6.1  软件测试的目的", 2)
    para(doc, "验证各模块在正常与边界输入下的正确性，确认约束检查（载重、时间窗、客户覆盖）可靠，"
              "并通过批量实验评估算法效果与稳定性。")
    heading(doc, "6.2  测试方法", 2)
    para(doc, "单元测试针对解析、评价、约束检查；集成测试覆盖「算例 → 求解 → 落库 → 报表」全链路；"
              "性能测试记录不同规模的求解耗时；算法效果采用多算例多次重复实验，"
              "并用配对 Wilcoxon 符号秩检验判断方法间差异是否显著。")
    heading(doc, "6.3  模块测试", 2)
    para(doc, f"（此处写 {name} 负责模块的测试用例表：编号、输入、预期、实际、结论。）")
    tip(doc, "至少 6 条用例，含 2 条异常输入；可用 store.py / run_experiments.py 的输出作为实测证据。")
    heading(doc, "6.4  系统总结", 2)
    para(doc, "（总结测试结论：功能是否完整、约束是否可靠、算法效果与耗时是否可接受、遗留问题。）")

    # ---------------- 第七章 ----------------
    heading(doc, "第七章  个人工作总结", 1)
    heading(doc, "1  个人工作总结", 2)
    para(doc, f"（{name}：写本人承担的工作、完成情况、遇到的问题与解决方式、收获。约 300–500 字。）")
    weeks = ["第一周", "第二周", "第三周", "第四周", "第五周", "第六周",
             "第七周", "第八周", "第九周", "第十周", "第十一周", "第十二周"]
    for i, w in enumerate(weeks, start=2):
        heading(doc, f"{i}  {w}周记", 2)
        para(doc, "（粘贴该周周记内容，约 100 字；对应文件：docs/周记/第NN周周记-*.docx）", size=11)

    # ---------------- 参考文献 ----------------
    heading(doc, "参考文献", 1)
    for ref in [
        "[1] 王宇春, 李雪梅, 等. 数据库系统及应用教程——SQL Server 2019[M]. 北京: 清华大学出版社, 2022.",
        "[2] 司守奎, 孙玺菁. 数学建模算法与应用[M]. 北京: 国防工业出版社, 2021.",
        "[3] 豆包. https://www.doubao.com/chat/.",
        "[4] DeepSeek. https://www.deepseek.com/.",
        "[5] GitHub Copilot. https://github.com/features/copilot.",
        "[6] Solomon M M. Algorithms for the vehicle routing and scheduling problems with time window constraints[J]. Operations Research, 1987, 35(2): 254-265.",
    ]:
        para(doc, ref, size=11, indent=False, space_after=2)

    return doc


def main() -> None:
    from report_base import base_doc

    OUT.mkdir(parents=True, exist_ok=True)
    for name, info in PEOPLE.items():
        path = OUT / f"报告骨架-{name}-{info['sid']}.docx"
        try:
            build(name, doc=base_doc(name, info["sid"])).save(str(path))
        except PermissionError:
            print(f"[跳过] {path.name} 正被 Word 占用，请关闭后重跑")
            continue
        d = Document(str(path))
        n_par = len([p for p in d.paragraphs if p.text.strip()])
        n_head = len([p for p in d.paragraphs if p.style.name.startswith("Heading")])
        print(f"[OK] {path.name}：{n_par} 段 / {n_head} 个标题 / {path.stat().st_size} bytes")


if __name__ == "__main__":
    main()
