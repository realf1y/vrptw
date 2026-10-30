# -*- coding: utf-8 -*-
"""报告附加件：自动目录 + 页脚页码 + 目录样式 + 封面修正"""
import copy

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def _set_font(run, size: float = 10.5, cn: str = "宋体") -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = False
    run.font.italic = False
    _r = run._element.get_or_add_rPr()
    _r.rFonts.set(qn("w:eastAsia"), cn)
    for tag in ("w:effect", "w:shadow", "w:outline", "w:emboss", "w:imprint"):
        el = _r.find(qn(tag))
        if el is not None:
            _r.remove(el)


def _field(par, instr: str, placeholder: str = "", size: float = 10.5) -> None:
    def mk():
        r = par.add_run()
        _set_font(r, size)
        return r

    r1 = mk()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    r1._r.append(fc)
    r2 = mk()
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r2._r.append(it)
    r3 = mk()
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate"); r3._r.append(fc)
    if placeholder:
        _set_font(par.add_run(placeholder), size)
    r4 = mk()
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); r4._r.append(fc)


def add_toc(doc) -> None:
    """在「课程任务及要求」前插入目录页"""
    anchor = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("课程任务及要求"):
            anchor = p
            break
    if anchor is None:
        return
    title = anchor.insert_paragraph_before("目　录")
    title.paragraph_format.page_break_before = True   # 段前分页，不产生空段落
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.runs[0]
    _set_font(r, 16, "黑体")
    r.font.bold = True

    toc = anchor.insert_paragraph_before("")
    _field(toc, ' TOC \\o "1-3" \\h \\z \\u ', "（打开 Word 后按 F9 更新目录）", 12)

    anchor.paragraph_format.page_break_before = True  # 正文另起一页，同样不用空段落


def fix_toc_styles(doc) -> None:
    """模板的 toc 1 是 22pt，太大；统一成 12pt 宋体"""
    for name in ("toc 1", "toc 2", "toc 3"):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.size = Pt(12)
        st.font.name = "Times New Roman"
        st.font.bold = False
        rPr = st.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:eastAsia"), "宋体")


def _fill_footer(footer, doc) -> None:
    footer.is_linked_to_previous = False
    # 整个页脚部件清空（模板里旧页码域可能藏在表格/内容控件里，只清段落清不掉）
    el = footer._element
    for child in list(el):
        el.remove(child)
    el.append(OxmlElement("w:p"))
    p = footer.paragraphs[0]
    p.style = doc.styles["Normal"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.first_line_indent = Pt(0)
    pf.left_indent = Pt(0)
    pf.line_spacing = 1.0
    _field(p, " PAGE ", "1", 10.5)


def _zero_indent(p_el) -> None:
    pPr = p_el.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    for a in ("w:firstLine", "w:firstLineChars", "w:hanging", "w:hangingChars"):
        ind.set(qn(a), "0")


def fix_cover(doc) -> None:
    """封面：补首页页眉（开了“首页不同”后首页页眉会变空）+ 取消首行缩进"""
    for section in doc.sections:
        if not section.different_first_page_header_footer:
            continue
        src = section.header
        dst = section.first_page_header
        el = dst._element
        for child in list(el):
            el.remove(child)
        for child in src._element:
            el.append(copy.deepcopy(child))

    # 封面：只清“有文字”的段落（标题、日期）；空段落保留模板的排版缩进
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            txt = "".join(n.text or "" for n in child.iter(qn("w:t"))).strip()
            if txt.startswith("课程任务及要求"):
                break
            if txt:
                _zero_indent(child)


def normalize_cover(doc) -> None:
    """封面表格：标签去掉字间空隙；所有文字统一 14pt 黑体（模板里空单元格填的值会漏样式）"""
    if not doc.tables:
        return
    t = doc.tables[0]

    # 1) 「题　　目：」拆成 3 个 run（题 / 空格 / 目：）→ 合成「题目：」
    c = t.rows[0].cells[0]
    par = c.paragraphs[0]
    for r in list(par.runs):
        r._element.getparent().remove(r._element)
    r0 = par.add_run("题目：")
    r0.font.size = Pt(14)
    r0.font.name = "黑体"
    r0._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

    # 2) 全表统一字号/字体/颜色；标签列左对齐、填写值居中（与模板一致）
    for row in t.rows:
        for ci, cell in enumerate(row.cells):
            is_label = ci in (0, 2)
            for p in cell.paragraphs:
                p.alignment = None if is_label else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(14)
                    r.font.name = "黑体"
                    r.font.bold = None
                    _r = r._element.get_or_add_rPr()
                    _r.rFonts.set(qn("w:eastAsia"), "黑体")
                    r.font.color.rgb = RGBColor(0, 0, 0)


def disable_auto_update(doc) -> None:
    """关闭“打开时自动更新所有域”（目录等不自动重排）"""
    st = doc.settings.element
    el = st.find(qn("w:updateFields"))
    if el is not None:
        st.remove(el)


def add_page_numbers(doc) -> None:
    """页脚只放一个居中页码；封面不显示；全篇连续编号

    模板里有两处坑：
    1) 页脚本身带页码域，只删文字会残留 → 重影；
    2) 第一章前有分节符，第 2 节带 pgNumType start=1 → 页码从第一章重新数，
       而且第 2 节还开了“首页不同”却没给首页页脚 → 那一页没页码。
    """
    for idx, section in enumerate(doc.sections):
        section.different_first_page_header_footer = (idx == 0)

        sp = section._sectPr
        pg = sp.find(qn("w:pgNumType"))
        if pg is not None:
            sp.remove(pg)
        if idx == 0:
            # 首页（封面）计为 0 且不显示 → 目录页正好是 1
            pg = OxmlElement("w:pgNumType")
            pg.set(qn("w:start"), "0")
            sp.append(pg)

        _fill_footer(section.footer, doc)
        try:
            _fill_footer(section.even_page_footer, doc)
        except Exception:
            pass
