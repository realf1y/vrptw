# -*- coding: utf-8 -*-
"""把 docs/截图 下的运行截图插到对应小节（幂等：已存在则跳过）"""
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

SHOT = Path("D:/vrptw/docs/截图")

PLAN = {
    "任绪锋": [
        ("168 个标准算例", "5-4_数据库.png",
         "图 5-1  实验记录写入 MySQL 数据库（Navicat 查看 experiment 表）"),
        ("改进版求得总里程 829.01", "5-1_算法输出.png",
         "图 5-2  算法在 C101 上的求解结果（命令行输出）"),
    ],
    "苍韵同": [
        ("优化前（最近邻）为 1870.73", "5-5_路线图.png",
         "图 5-1  路线图页面（可按按钮切换方法）"),
        ("600 次运行", "5-2_批量实验.png",
         "图 5-2  批量实验执行过程与汇总输出"),
        ("60 行 KPI 汇总", "5-3_报表页.png",
         "图 5-3  实验报表页（KPI 卡片与逐算例对比）"),
    ],
}


def _insert_after(par, img: Path, caption: str, width_cm: float = 15.0) -> None:
    new_p = OxmlElement("w:p")
    par._p.addnext(new_p)
    np = Paragraph(new_p, par._parent)
    np.alignment = WD_ALIGN_PARAGRAPH.CENTER
    np.paragraph_format.line_spacing = 1.0
    np.add_run().add_picture(str(img), width=Cm(width_cm))

    cap_p = OxmlElement("w:p")
    new_p.addnext(cap_p)
    cp = Paragraph(cap_p, par._parent)
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.add_run(caption).font.size = Pt(10.5)


def add_screenshots(doc, name: str) -> int:
    items = PLAN.get(name, [])
    if not items:
        return 0
    existing = "".join(p.text for p in doc.paragraphs)
    n = 0
    for anchor, img, caption in items:
        if caption in existing:
            continue
        target = next((p for p in doc.paragraphs if anchor in p.text), None)
        if target is None:
            print(f"[截图] 未找到锚点：{anchor}")
            continue
        img_path = SHOT / img
        if not img_path.exists():
            print(f"[截图] 缺文件：{img_path.name}")
            continue
        _insert_after(target, img_path, caption)
        n += 1
    return n
