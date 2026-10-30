# -*- coding: utf-8 -*-
"""格式自检：生成/修改完报告后跑一遍，确认符合固化规范

用法: python check_format.py [文件...]   不带参数则检查 docs 下所有 报告-最终-*.docx
"""
import glob
import sys

from docx import Document
from docx.oxml.ns import qn


def check(path):
    d = Document(path)
    issues = []

    # 1) 封面表：标签（以：结尾）左对齐，值居中
    if d.tables:
        seen = set()
        for row in d.tables[0].rows:
            for c in row.cells:
                if id(c._tc) in seen:
                    continue
                seen.add(id(c._tc))
                txt = c.text.strip()
                if not txt:
                    continue
                a = c.paragraphs[0].alignment
                if txt.endswith("：") and a is not None:
                    issues.append(f"封面标签未左对齐：{txt[:8]}")
                if not txt.endswith("：") and a is None:
                    issues.append(f"封面值未居中：{txt[:8]}")

    # 2) 彩色文字
    colored = 0
    for p in d.paragraphs:
        for r in p.runs:
            if r.font.color and r.font.color.rgb is not None and str(r.font.color.rgb) != "000000":
                colored += 1
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        if r.font.color and r.font.color.rgb is not None and str(r.font.color.rgb) != "000000":
                            colored += 1
    if colored:
        issues.append(f"彩色文字 {colored} 处")

    # 3) 表格首行缩进清零
    bad = 0
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    pPr = p._p.pPr
                    ind = pPr.find(qn("w:ind")) if pPr is not None else None
                    if ind is None or ind.get(qn("w:firstLine")) != "0" or ind.get(qn("w:firstLineChars")) != "0":
                        bad += 1
    if bad:
        issues.append(f"表格未清零缩进 {bad} 处")

    # 4) 图片段落 1.0 行距
    bad = 0
    for p in d.paragraphs:
        if p._p.findall(".//" + qn("w:drawing")):
            if abs((p.paragraph_format.line_spacing or 0) - 1.0) > 1e-6:
                bad += 1
    if bad:
        issues.append(f"图片段落非 1.0 行距 {bad} 处")

    # 5) 目录
    if "TOC" not in d.element.xml:
        issues.append("缺目录域")
    if not any(p.text.strip().startswith("目") and p.paragraph_format.page_break_before
               for p in d.paragraphs):
        issues.append("目录未设段前分页")

    # 6) 页脚：每节默认页脚恰好 1 个 PAGE 域、无“第/页”、首页页脚为空
    for i, s in enumerate(d.sections):
        p = s.footer.paragraphs[0]
        if p._p.xml.count('w:fldCharType="begin"') != 1:
            issues.append(f"第{i}节页脚域数不为 1")
        if "第" in p.text or "页" in p.text:
            issues.append(f"第{i}节页脚含“第/页”")
        if s.first_page_footer.paragraphs[0].text.strip():
            issues.append(f"第{i}节首页页脚非空")

    # 7) 自动更新已关
    if d.settings.element.find(qn("w:updateFields")) is not None:
        issues.append("updateFields 未关闭")
    if "w:dirty" in d.element.xml:
        issues.append("仍存在 w:dirty（会触发自动更新）")

    return issues


files = sys.argv[1:] or [f for f in sorted(glob.glob("D:/vrptw/docs/报告-最终-*.docx"))
                           if ".old." not in f]
ok = 0
for f in files:
    iss = check(f)
    name = f.replace("\\", "/").split("/")[-1]
    print(("✅ PASS  " if not iss else "❌ FAIL  ") + name)
    for i in iss[:8]:
        print("     -", i)
    ok += (not iss)
print(f"\n结果：{ok}/{len(files)} 通过")
